"""LinkedIn and spreadsheet import, Excel export."""
import csv
import io
import json

import openpyxl

import db

FIELDS = ("name", "first_name", "last_name", "company", "title", "email", "linkedin_url", "how_connected", "notes")
LINKEDIN_MAPPING = {"First Name": "first_name", "Last Name": "last_name", "Email Address": "email",
                    "Company": "company", "Position": "title", "URL": "linkedin_url"}


def import_rows(con, rows, mapping, source, how_connected=""):
    report = {"added": 0, "skipped": 0, "failed": []}
    for n, row in enumerate(rows, start=1):
        f = {field: str(row.get(col) or "").strip() for col, field in mapping.items() if field}
        first, last = f.pop("first_name", ""), f.pop("last_name", "")
        name = f.pop("name", "") or f"{first} {last}"
        if not db.clean_name(name):
            report["failed"].append(f"row {n}: missing name")
            continue
        f["how_connected"] = f.get("how_connected") or how_connected
        if db.add_contact(con, name, source, **f) is None:
            report["skipped"] += 1
        else:
            report["added"] += 1
    return report


def import_linkedin(con, data):
    lines = data.decode("utf-8-sig").splitlines()
    start = next((i for i, line in enumerate(lines) if line.startswith("First Name")), None)
    if start is None:
        raise ValueError("This is not a LinkedIn Connections.csv (no 'First Name' header found).")
    return import_rows(con, list(csv.DictReader(lines[start:])), LINKEDIN_MAPPING,
                       "linkedin_csv", "1st-degree LinkedIn connection")


def read_table(filename, data, sheet=None):
    if filename.lower().endswith(".csv"):
        sheets, grid = ["CSV"], list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    else:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise ValueError(f"Could not read {filename}. Upload an .xlsx or .csv file.") from e
        sheets = wb.sheetnames
        ws = wb[sheet] if sheet in sheets else wb[sheets[0]]
        grid = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    # Exports such as PitchBook put title rows above the header: take the first row with 2+ filled cells.
    start = next((i for i, r in enumerate(grid) if sum(1 for v in r if v.strip()) >= 2), None)
    if start is None:
        return sheets, [], []
    headers = [h.strip() for h in grid[start]]
    rows = [dict(zip(headers, r)) for r in grid[start + 1:] if any(v.strip() for v in r)]
    return sheets, headers, rows


def _signature(headers):
    return "|".join(h.lower() for h in headers)


def saved_mapping(con, headers):
    row = con.execute("SELECT mapping FROM column_mappings WHERE signature = ?", (_signature(headers),)).fetchone()
    return json.loads(row["mapping"]) if row else {}


def save_mapping(con, headers, mapping):
    con.execute("INSERT OR REPLACE INTO column_mappings (signature, mapping) VALUES (?, ?)",
                (_signature(headers), json.dumps(mapping)))


def export_xlsx(con):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for table in ("contacts", "companies", "messages", "applications", "email_templates", "profile"):
        cur = con.execute(f"SELECT * FROM {table}")
        ws = wb.create_sheet(table)
        ws.append([d[0] for d in cur.description])
        for row in cur:
            ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
