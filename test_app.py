import datetime as dt
import io

import openpyxl
import pytest

import ai
import db
import importers
import outlook


@pytest.fixture
def con(tmp_path):
    c = db.connect(tmp_path / "test.db")
    yield c
    c.close()


def sent_contact(con, sent="2026-10-01T09:00:00Z"):
    cid = db.add_contact(con, "Jane Doe", "paste", company="Stripe")
    mid = db.record_draft(con, cid, "initial", "Hi", "Body")
    db.mark_sent(con, mid, sent)
    return cid, mid


# --- db ---

def test_duplicate_contact_same_name_and_company_is_skipped(con):
    assert db.add_contact(con, "Jane Doe", "paste", company="Stripe")
    assert db.add_contact(con, "  jane   doe ", "paste", company="stripe") is None
    assert db.add_contact(con, "Jane Doe", "paste", company="Figma")
    assert db.add_contact(con, "Jane Doe", "paste")
    assert db.add_contact(con, "jane doe", "paste") is None


def test_add_contact_requires_name(con):
    with pytest.raises(ValueError):
        db.add_contact(con, "   ", "paste")


def test_draft_then_sent_sets_followup_date(con):
    cid = db.add_contact(con, "Jane Doe", "paste", company="Stripe")
    mid = db.record_draft(con, cid, "initial", "Hi", "Body")
    assert db.get_contact(con, cid)["status"] == "drafted"
    db.mark_sent(con, mid, "2026-10-01T09:00:00Z")
    c = db.get_contact(con, cid)
    assert (c["status"], c["next_followup_date"], c["followup_count"]) == ("sent", "2026-10-06", 0)


def test_followup_days_setting(con):
    con.execute("UPDATE profile SET followup_days = 7")
    cid, _ = sent_contact(con)
    assert db.get_contact(con, cid)["next_followup_date"] == "2026-10-08"


def test_two_sent_followups_count_and_stay_due(con):
    cid, _ = sent_contact(con)
    for sent in ("2026-10-06T09:00:00Z", "2026-10-11T09:00:00Z"):
        mid = db.record_draft(con, cid, "followup", "Re: Hi", "Nudge")
        db.mark_sent(con, mid, sent)
    c = db.get_contact(con, cid)
    assert (c["followup_count"], c["next_followup_date"]) == (2, "2026-10-16")
    due = db.today_view(con, dt.date(2026, 10, 16))["due"]
    assert [r["id"] for r in due] == [cid]
    assert due[0]["followup_count"] >= 2  # template shows "suggest dead"


def test_mark_sent_twice_is_noop(con):
    cid, mid = sent_contact(con)
    db.mark_sent(con, mid, "2026-10-09T09:00:00Z")
    assert db.get_contact(con, cid)["next_followup_date"] == "2026-10-06"


def test_reply_clears_followup(con):
    cid, mid = sent_contact(con)
    db.mark_replied(con, mid, "2026-10-02T10:00:00Z")
    c = db.get_contact(con, cid)
    assert (c["status"], c["next_followup_date"]) == ("replied", None)
    assert db.today_view(con, dt.date(2026, 10, 30))["due"] == []


def test_manual_status_overrides_automation(con):
    cid, mid = sent_contact(con)
    db.set_status(con, cid, "call")
    assert db.get_contact(con, cid)["next_followup_date"] is None
    db.mark_replied(con, mid, "2026-10-02T10:00:00Z")
    assert db.get_contact(con, cid)["status"] == "call"


def test_today_lists_old_unsent_drafts_and_stale_apps(con):
    cid = db.add_contact(con, "Jane Doe", "paste", company="Stripe")
    db.record_draft(con, cid, "initial", "Hi", "Body")
    con.execute("UPDATE messages SET created_at = '2026-10-01 12:00:00'")
    aid = db.add_application(con, "Stripe", "PM Intern")
    con.execute("UPDATE applications SET updated_at = '2026-09-01 12:00:00'")
    assert db.today_view(con, dt.date(2026, 10, 2))["unsent"] == []
    view = db.today_view(con, dt.date(2026, 10, 3))
    assert len(view["unsent"]) == 1
    assert [a["id"] for a in view["stale_apps"]] == [aid]


def test_application_status_change_sets_applied_date(con):
    aid = db.add_application(con, "Stripe", "PM Intern")
    con.execute("UPDATE applications SET updated_at = '2026-09-01 12:00:00'")
    db.update_application(con, aid, "applied")
    a = con.execute("SELECT * FROM applications WHERE id = ?", (aid,)).fetchone()
    assert a["applied_date"] and a["updated_at"] > "2026-09-01 12:00:00"


# --- ai ---

def test_fill_data_leaves_unknown_and_empty_placeholders_visible():
    out = ai.fill_data("Hi {first_name} at {company}, {mystery}", {"first_name": "Jane", "company": ""})
    assert out == "Hi Jane at {company}, {mystery}"
    assert ai.unfilled(out + " [[hook: x]]") == ["{company}", "{mystery}", "[[hook: x]]"]


def test_slots_extract_and_apply_keeps_missing_marker():
    text = "A [[hook: why them]] B [[why_company: why\nStripe]]"
    assert ai.extract_slots(text) == {"hook": "why them", "why_company": "why\nStripe"}
    assert ai.apply_slots(text, {"hook": " We both did CS. "}) == "A We both did CS. B [[why_company: why\nStripe]]"


def test_parse_json_handles_fences_and_prose():
    assert ai.parse_json('Sure!\n```json\n{"a": [1]}\n```') == {"a": [1]}
    with pytest.raises(ai.AIError):
        ai.parse_json("no json here")
    with pytest.raises(ai.AIError):
        ai.parse_json("{not: valid}")


PROFILE = {"name": "Alex Lee", "program": "CS", "target_roles": "Product, eng", "voice_samples": "hey!"}


def test_draft_fills_data_first_then_only_slots(monkeypatch):
    seen = {}

    def fake_fill(slots, context, voice):
        seen.update(slots)
        assert voice == "hey!" and "Stripe" in context
        return {"hook": "Loved your talk."}

    monkeypatch.setattr(ai, "fill_slots", fake_fill)
    tpl = {"subject": "{target_role} at {company}", "body": "Hi {first_name},\n[[hook: about {company}]]\n{my_name}"}
    contact = {"name": "Cher", "title": "", "company": "Stripe"}
    assert ai.draft(tpl, PROFILE, contact, None, []) == ("Product at Stripe", "Hi Cher,\nLoved your talk.\nAlex Lee")
    assert seen == {"hook": "about Stripe"}


def test_draft_without_slots_makes_no_ai_call(monkeypatch):
    monkeypatch.setattr(ai, "fill_slots", lambda *a: pytest.fail("should not call the AI"))
    tpl = {"subject": "Hi", "body": "Hi {first_name}"}
    assert ai.draft(tpl, PROFILE, {"name": "Jane Doe"}, None, []) == ("Hi", "Hi Jane")



# --- importers ---

LINKEDIN_CSV = (
    "﻿Notes:\n"
    "\"When exporting your connection data, you may notice that some of the email addresses are missing.\"\n"
    "\n"
    "First Name,Last Name,URL,Email Address,Company,Position,Connected On\n"
    "Jane,Doe,https://www.linkedin.com/in/janedoe,,Stripe,PM,01 Oct 2026\n"
    "Bob,,https://www.linkedin.com/in/bob,,,,02 Oct 2026\n"
    ",,,,,,\n"
    "Jane,Doe,https://www.linkedin.com/in/janedoe,,Stripe,PM,01 Oct 2026\n"
).encode()


def test_linkedin_import_handles_bom_preamble_blanks_and_duplicates(con):
    report = importers.import_linkedin(con, LINKEDIN_CSV)
    assert report == {"added": 2, "skipped": 1, "failed": ["row 3: missing name"]}
    jane = db.list_contacts(con, source="linkedin_csv")[-1]
    assert (jane["name"], jane["company"], jane["title"]) == ("Jane Doe", "Stripe", "PM")
    assert jane["how_connected"] == "1st-degree LinkedIn connection"


def test_linkedin_import_rejects_other_csv(con):
    with pytest.raises(ValueError):
        importers.import_linkedin(con, b"a,b\n1,2\n")


def test_excel_import_skips_title_rows_and_remembers_mapping(con):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["PitchBook search results"])
    ws.append([])
    ws.append(["Full Name", "Employer", "Job Title"])
    ws.append(["Ann Lee", "Figma", "Recruiter"])
    ws.append([None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    sheets, headers, rows = importers.read_table("people.xlsx", buf.getvalue())
    assert headers == ["Full Name", "Employer", "Job Title"]
    assert rows == [{"Full Name": "Ann Lee", "Employer": "Figma", "Job Title": "Recruiter"}]
    mapping = {"Full Name": "name", "Employer": "company", "Job Title": "title"}
    importers.save_mapping(con, headers, mapping)
    assert importers.saved_mapping(con, headers) == mapping
    assert importers.import_rows(con, rows, mapping, "import")["added"] == 1


def test_read_table_rejects_garbage():
    with pytest.raises(ValueError):
        importers.read_table("people.xlsx", b"not an excel file")


def test_export_has_a_sheet_per_table(con):
    db.add_contact(con, "Jane Doe", "paste")
    wb = openpyxl.load_workbook(io.BytesIO(importers.export_xlsx(con)))
    assert {"contacts", "companies", "messages", "applications"} <= set(wb.sheetnames)
    assert wb["contacts"]["B2"].value == "Jane Doe"


# --- outlook ---

def graph_msg(addr, when, draft=False):
    return {"isDraft": draft, "from": {"emailAddress": {"address": addr}}, "sentDateTime": when, "receivedDateTime": when}


def test_classify_conversation():
    me = {"me@uwaterloo.ca"}
    assert outlook.classify([graph_msg("me@uwaterloo.ca", "2026-10-01T09:00:00Z", draft=True)], me, "2026-10-01 08:00:00") == (None, None)
    convo = [graph_msg("Me@UWaterloo.ca", "2026-10-01T09:00:00Z"), graph_msg("jane@stripe.com", "2026-10-02T10:00:00Z")]
    assert outlook.classify(convo, me, "2026-10-01 08:00:00") == ("2026-10-01T09:00:00Z", "2026-10-02T10:00:00Z")
    # A follow-up drafted later in the same thread ignores the earlier send and reply.
    assert outlook.classify(convo, me, "2026-10-05 08:00:00") == (None, None)


def test_draft_and_mailto_without_recipient():
    assert outlook.draft_payload("", "S", "B")["toRecipients"] == []
    assert outlook.draft_payload("a@b.com", "S", "B")["toRecipients"] == [{"emailAddress": {"address": "a@b.com"}}]
    assert outlook.mailto_link("", "Hi there", "Line 1\nLine 2") == "mailto:?subject=Hi%20there&body=Line%201%0ALine%202"


# --- app ---

def test_every_page_renders_and_contact_flow_works(tmp_path, monkeypatch):
    import app as webapp
    monkeypatch.setattr(outlook, "enabled", lambda: False)
    webapp.app.config["DB_PATH"] = tmp_path / "web.db"
    client = webapp.app.test_client()
    for path in ["/", "/contacts", "/companies", "/find", "/applications", "/templates", "/profile", "/import"]:
        assert client.get(path).status_code == 200, path
    r = client.post("/contacts", data={"name": "Jane Doe", "company": "Stripe", "source": "paste"})
    assert r.status_code == 302
    assert b"Jane Doe" in client.get(r.headers["Location"]).data
    assert client.post("/contacts", data={"name": "jane doe", "company": "stripe", "source": "paste"}).status_code == 302
    assert client.get("/companies/1").status_code == 200
    r = client.post("/contacts/1/save_draft", data={"subject": "Hi", "body": "Hello", "kind": "initial", "action": "local"})
    assert r.status_code == 302
    page = client.get("/contacts/1").data
    assert b"mailto:" in page and b"No email on file" in page


def test_missing_api_key_becomes_readable_ai_error(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.setattr(ai, "_client", None)
    with pytest.raises(ai.AIError, match="OPENAI_API_KEY"):
        ai.ask_json("system", "prompt")


def test_call_uses_responses_api_with_web_search_only_when_asked(monkeypatch):
    calls = []

    class FakeResponses:
        @staticmethod
        def create(**kwargs):
            calls.append(kwargs)
            return type("R", (), {"output_text": '{"ok": true}'})()

    monkeypatch.setattr(ai, "client", lambda: type("C", (), {"responses": FakeResponses})())
    assert ai.ask_json("be brief", "hi", search=True) == {"ok": True}
    assert ai.ask_json("be brief", "hi") == {"ok": True}
    assert calls[0]["model"] == ai.MODEL and calls[0]["instructions"] == "be brief" and calls[0]["input"] == "hi"
    assert calls[0]["tools"] == [{"type": "web_search"}] and "tools" not in calls[1]


# --- final review fixes ---

def test_classify_treats_any_own_alias_as_mine():
    me = {"s2kim@uwaterloo.ca", "sean.kim@uwaterloo.ca"}
    convo = [graph_msg("Sean.Kim@uwaterloo.ca", "2026-10-01T09:00:00Z")]
    assert outlook.classify(convo, me, "2026-10-01 08:00:00") == ("2026-10-01T09:00:00Z", None)


def test_my_addresses_include_proxy_aliases(monkeypatch):
    monkeypatch.setattr(outlook, "_me", None)
    monkeypatch.setattr(outlook, "_graph", lambda *a, **k: {
        "mail": "Sean.Kim@uwaterloo.ca", "userPrincipalName": "s2kim@uwaterloo.ca",
        "proxyAddresses": ["SMTP:sean.kim@uwaterloo.ca", "smtp:s2kim@edu.uwaterloo.ca"]})
    assert outlook._my_addresses() == {"sean.kim@uwaterloo.ca", "s2kim@uwaterloo.ca", "s2kim@edu.uwaterloo.ca"}


def test_is_connected_false_when_token_cannot_refresh(monkeypatch):
    class FakeMsal:
        def get_accounts(self):
            return [{"username": "me"}]

        def acquire_token_silent(self, scopes, account):
            return {"error": "invalid_grant"}

    monkeypatch.setattr(outlook, "_msal", lambda: FakeMsal())
    assert outlook.is_connected() is False


def test_network_failure_becomes_outlook_error(monkeypatch):
    import requests

    def offline(*a, **k):
        raise requests.ConnectionError("no network")

    monkeypatch.setattr(outlook, "_token", lambda: "t")
    monkeypatch.setattr(outlook.requests, "request", offline)
    with pytest.raises(outlook.OutlookError):
        outlook._graph("GET", "/me")


def test_sync_only_polls_open_contacts_and_applies_results(con, monkeypatch):
    live = db.add_contact(con, "Jane Doe", "paste", company="Stripe")
    gone = db.add_contact(con, "Bob Roe", "paste", company="Stripe")
    m1 = db.record_draft(con, live, "initial", "Hi", "Body", "id1", "conv1")
    db.record_draft(con, gone, "initial", "Hi", "Body", "id2", "conv2")
    db.set_status(con, gone, "dead")
    con.execute("UPDATE messages SET created_at = '2026-10-01 08:00:00'")
    polled = []

    def fake_conversation(cid):
        polled.append(cid)
        return [graph_msg("me@uwaterloo.ca", "2026-10-01T09:00:00Z"),
                graph_msg("jane@stripe.com", "2026-10-02T10:00:00Z")]

    monkeypatch.setattr(outlook, "_conversation", fake_conversation)
    monkeypatch.setattr(outlook, "_my_addresses", lambda: {"me@uwaterloo.ca"})
    assert outlook.sync(con, force=True) is True
    assert polled == ["conv1"]
    msg = con.execute("SELECT sent_at, replied_at FROM messages WHERE id = ?", (m1,)).fetchone()
    assert (msg["sent_at"], msg["replied_at"]) == ("2026-10-01T09:00:00Z", "2026-10-02T10:00:00Z")
    assert db.get_contact(con, live)["status"] == "replied"


def test_scopes_can_read_profile_but_never_send():
    assert "User.Read" in outlook.SCOPES and "Mail.Send" not in outlook.SCOPES


def test_cross_site_posts_are_rejected(tmp_path, monkeypatch):
    import app as webapp
    monkeypatch.setattr(outlook, "enabled", lambda: False)
    webapp.app.config["DB_PATH"] = tmp_path / "web.db"
    client = webapp.app.test_client()
    data = {"name": "X", "source": "paste"}
    assert client.post("/contacts", data=data, headers={"Sec-Fetch-Site": "cross-site"}).status_code == 403
    assert client.post("/contacts", data=data, headers={"Origin": "https://evil.example"}).status_code == 403
    assert client.post("/contacts", data=data, headers={"Origin": "http://localhost", "Sec-Fetch-Site": "same-origin"}).status_code == 302


def test_safe_url_blocks_script_links():
    import app as webapp
    f = webapp.app.jinja_env.filters["safe_url"]
    assert f("javascript:alert(1)") == "#"
    assert f(" JaVaScRiPt:alert(1)") == "#"
    assert f("https://www.linkedin.com/in/x") == "https://www.linkedin.com/in/x"
    assert f("") == "#"
